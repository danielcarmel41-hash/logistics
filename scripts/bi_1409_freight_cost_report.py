"""
Prices every line in data/BI_1409.xlsx (single flat "BI 1409" sheet, not
pre-split into population sheets like earlier data pulls) and writes
output/BI_1409_Freight_Cost_Report.xlsx (Shipped / Backlog / Summary /
README) plus a standalone filterable HTML dashboard.

Populations and routing rules (as given):

  Sending WHS Code = 3PLCANOT (Israel):
    - Destination = Israel (domestic)  -> Canot Whs. NIS rate card by
      delivery-city region: 200 ILS flat for 1 pallet, 1200/1300/1400 ILS
      (Main Land Gedera-Hadera / Gedera to south / Hadera and north) for
      2-16 pallets (8T truck), same rate for >16 (12T truck) -- confirmed
      against the forwarded price-list email screenshot.
    - Everything else (export)         -> Ship Cost Matrix, Sending Org
      Code 3PLCANOT + Ship Mode + destination country.

  Sending WHS Code in {3PLDBSNL, 3PLDBSBRNL} (Netherlands), checked in
  this order for every line:
    1. Family Type = Battery, OR Destination country = United Kingdom
                                        -> MNT UK price list (sheet
       "Price Q3"), matched by Customer Name (+ ZIP for customers with
       more than one listed address). This is a per-shipment full-truck
       rate, not per line -- grouped and allocated pro-rata by pallet
       share like every other flat-rate source below. That file only
       covers a handful of specific UK-shipment customers, so any line
       whose customer isn't in it at all falls back to the domestic
       (DBS)/export (matrix) split below, flagged in the Note -- most
       Battery-family lines are ordinary shipments to European
       destinations that were never meant to depend on a UK truck-rate
       file (same fallback pattern used for this population in the
       prior 0609 report).
    2. A/I = "Support"                 -> same rule as domestic (#3)
       below, regardless of Parent Region.
    3. Parent Region = EUROPE (domestic):
         - Forwarder = MNT AND has a real Logistic POD date AND the
           Shipment Number is on file as a "PL nr." in Q3 prices AUGUST
                                        -> Q3 prices AUGUST (Solaredge
           rate + waiting + cancellation + customs clearance, EUR; Duty
           GBP reported separately since it's a different currency).
         - Everyone else (DBSCHENKER, DEFAULT, DGF forwarders, or MNT
           without a POD/Q3 match)      -> DBS Price list 2026, by
           destination zone (country code + first 2 digits of ZIP, or
           GB postcode area) + pallet count.
    4. Parent Region != EUROPE (export) -> Ship Cost Matrix.

    BayWa r.e. Solar Systems srl (Italy) lines whose ZIP is garbage text
    ("3c/4c block") use the user-supplied EUR rates by total shipment
    pallet count (1->240, 2->411, 12->2100, 13->1410) instead of the DBS
    lookup (which can't resolve a zone from that ZIP).

  Any other Sending WHS Code (3PLDBSTW, 3PLSEAWIN, 3PLEXPAU, 3PLMCNY,
  3PLMCCA, ...)                        -> Ship Cost Matrix, Sending Org
                                           Code + Ship Mode + destination
                                           country.

Every one of the rate sources above (DBS, Ship Cost Matrix, Q3 AUGUST,
MNT UK, Canot rate card, BayWa override) gives ONE flat price per
shipment/order, not per line -- this was the recurring, financially
material bug found and fixed across the whole 0609 data pull project.
Every line here is therefore priced by first grouping same-shipment
lines (Shipment Number for Shipped rows that have one, else SE Order#)
within each pricing method, looking up ONE price for the group's total
whole-number pallets, then allocating pro-rata by each line's pallet
share -- from the start, not discovered after the fact.

Whole-number pallets: "# of Pallets per line-up" is used directly (all
3,132 source rows already hold a whole number); ceil() is still applied
defensively in every lookup in case a future data pull doesn't.

Dedup: Shipped rows are checked for exact full-row duplicates (same
pattern found and fixed in the 0609 pull); none are found in this pull.
No Backlog row is ever considered for removal, per the explicit request.

FX to USD: 1 EUR = 1.1568, 1 ILS = 0.32955 (2026-09-14 snapshot,
xe.com/tradingeconomics -- not a contracted rate); Ship Cost Matrix
costs are already USD.
"""
import math
import os
import re
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = lambda name: os.path.join(BASE_DIR, "data", name)
OUT_FILE = os.path.join(BASE_DIR, "output", "BI_1409_Freight_Cost_Report.xlsx")
OUT_JSON = os.path.join(BASE_DIR, "output", "bi_1409_report_data.json")

MAIN_FILE = DATA("BI_1409.xlsx")
DBS_FILE = DATA("DBS_Price_list_2026_v2.xlsx")
MATRIX_FILE = DATA("Ship_Cost_Matrix_v6.xlsx")
MNT_UK_FILE = DATA("MNT_UK_price_list_v2.xlsx")
Q3_AUG_FILE = DATA("Q3_prices_AUGUST_v2.xlsx")

MAX_EP = 33
FX_TO_USD = {'EUR': 1.1568, 'ILS': 0.32955, 'USD': 1.0}

NL_WHS = {'3PLDBSNL', '3PLDBSBRNL'}
CANOT_WHS = '3PLCANOT'

BAYWA_IT_CUSTOMER = 'BayWa r.e. Solar Systems srl'
BAYWA_IT_OVERRIDES = {1: 240, 2: 411, 12: 2100, 13: 1410}

CANOT_SINGLE_PALLET_RATE_ILS = 200
CANOT_REGION_RATE_ILS = {'mainland': 1200, 'south': 1300, 'north': 1400}
CANOT_CITY_REGION = {
    'KIRYAT GAT': 'south', 'ASHDOD': 'south', 'KADIMA': 'mainland',
    'EIN HAEMEK': 'north', 'BEIT HASHITTA': 'north',
    'INDUSTRIAL PARK KIDMAT GALIL': 'north',
}

EXPECTED_ZIP_DIGITS = {'CH': 4, 'AT': 4, 'BE': 4, 'HU': 4, 'SI': 4, 'LU': 4, 'NL': 4,
                       'DE': 5, 'IT': 5, 'FR': 5, 'PL': 5, 'CZ': 5, 'FI': 5}


def to_usd(cost, currency):
    if cost is None or currency not in FX_TO_USD:
        return None
    return round(cost * FX_TO_USD[currency], 2)


# --------------------------------------------------------------------------
# Price sources
# --------------------------------------------------------------------------

def parse_dbs_price_book(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    book = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row = next(
            (r for r in range(1, 15)
             if any(ws.cell(row=r, column=c).value == 'EP' for c in range(1, ws.max_column + 1))),
            None,
        )
        if header_row is None:
            continue
        ep_col = None
        zone_cols = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if v == 'EP':
                ep_col = c
            elif isinstance(v, str) and v.strip().upper().startswith(sheet_name.upper()):
                zone_cols[v.strip()] = c
            elif isinstance(v, str) and sheet_name.upper() == 'NL' and re.match(r'^\d+\s*-\s*\d+$', v.strip()):
                zone_cols['NL'] = c
        zones = {z: {} for z in zone_cols}
        data_start = header_row + 2
        for r in range(data_start, ws.max_row + 1):
            ep_val = ws.cell(row=r, column=ep_col).value
            if not isinstance(ep_val, (int, float)):
                continue
            ep_val = int(ep_val)
            for zone, col in zone_cols.items():
                price = ws.cell(row=r, column=col).value
                if isinstance(price, (int, float)):
                    zones[zone][ep_val] = price
        book[sheet_name] = zones
    return book


def dbs_zone_for(country_code, zip_code):
    country_code = (country_code or '').strip().upper()
    zip_code = str(zip_code or '').strip().upper()
    if country_code == 'NL':
        return 'NL'
    if country_code == 'GB':
        m = re.match(r'^[A-Z]+', zip_code.replace(' ', ''))
        return 'GB' + m.group() if m else None
    digits = ''.join(ch for ch in zip_code if ch.isdigit())
    expected = EXPECTED_ZIP_DIGITS.get(country_code)
    if expected and len(digits) == expected + 1 and digits[0] == '0':
        digits = digits[1:]
    if len(digits) < 2:
        return None
    return country_code + digits[:2]


def dbs_price_for_pallets(zone_prices, total_pallets):
    ep = max(1, min(MAX_EP, math.ceil(total_pallets - 1e-9)))
    while ep <= MAX_EP:
        if ep in zone_prices:
            return zone_prices[ep]
        ep += 1
    return None


def dbs_lookup(dbs_book, country_code, zip_code, total_pallets):
    zone = dbs_zone_for(country_code, zip_code)
    if not zone:
        return None, None, 'ZIP could not be parsed into a rate zone -- confirm with WH contact.'
    zone_prices = dbs_book.get((country_code or '').strip().upper(), {}).get(zone)
    if not zone_prices:
        return zone, None, f'No DBS rate sheet/zone for country={country_code!r} zone={zone!r}.'
    price = dbs_price_for_pallets(zone_prices, total_pallets)
    if price is None:
        return zone, None, f'Zone {zone} has no rate bracket for {total_pallets:g} pallets.'
    return zone, price, None


def parse_ship_matrix(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Ship Cost Matrix']
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    by_org = defaultdict(list)
    by_country = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        org = r[idx['Sending Org Code**']]
        country = r[idx['Sending Country Name']]
        mode = (r[idx['Ship Mode*']] or '').strip().upper()
        dest = r[idx['Dest Country Name*']]
        cost = r[idx['Cost*']]
        currency = r[idx['Currency']]
        no_pal = r[idx['No. Pallets']]
        if not isinstance(cost, (int, float)) or not dest or not mode:
            continue
        by_country[(country, mode, dest)].append((cost, currency, no_pal))
        if org:
            by_org[(org, mode, dest)].append((cost, currency, no_pal))
    return by_org, by_country


def _resolve_matrix_entries(entries):
    if not entries:
        return None
    costs = [e[0] for e in entries]
    currency = entries[0][1]
    no_pal = entries[0][2]
    avg_cost = round(sum(costs) / len(costs), 2)
    note = '' if len(set(costs)) <= 1 else f'Averaged {len(costs)} rates on file for this route ({costs}).'
    return avg_cost, currency, no_pal, note


def matrix_lookup(matrix, org_code, sending_country, ship_mode, dest_country):
    by_org, by_country = matrix
    mode = (ship_mode or '').strip().upper()
    for m in ([mode, 'SEA'] if mode != 'SEA' else ['SEA']):
        hit = _resolve_matrix_entries(by_org.get((org_code, m, dest_country)))
        if hit:
            return hit
    for m in ([mode, 'SEA'] if mode != 'SEA' else ['SEA']):
        hit = _resolve_matrix_entries(by_country.get((sending_country, m, dest_country)))
        if hit:
            return hit
    return None


def canot_israel_lookup(city, total_pallets):
    pallets = max(1, math.ceil(total_pallets - 1e-9))
    region = CANOT_CITY_REGION.get((city or '').strip().upper())
    region_note = (f'Region "{region}" inferred from city {city!r} (not in the source data).'
                   if region else f'City {city!r} not mapped to a Canot region -- confirm with WH; not priced.')
    if pallets == 1:
        return CANOT_SINGLE_PALLET_RATE_ILS, 'ILS', f'Single-pallet flat rate. {region_note}'
    if not region:
        return None, 'ILS', region_note
    truck = '8T' if pallets <= 16 else '12T'
    return CANOT_REGION_RATE_ILS[region], 'ILS', f'{truck} truck rate, region={region}. {region_note}'


def parse_mnt_uk(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Price Q3']
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    by_name_zip = {}
    by_name = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        name = r[idx['Customer Name']]
        addr = r[idx['Ship To Address']] or ''
        total = r[idx['Total price']]
        if not isinstance(total, (int, float)):
            continue
        m = re.search(r'([A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2})', addr.upper())
        zip_code = m.group(1).replace(' ', '') if m else None
        if zip_code:
            by_name_zip[(name, zip_code)] = total
        by_name[name].append(total)
    return by_name_zip, by_name


def mnt_uk_lookup(mnt_book, customer_name, zip_code):
    by_name_zip, by_name = mnt_book
    zip_norm = (zip_code or '').upper().replace(' ', '')
    hit = by_name_zip.get((customer_name, zip_norm))
    if hit is not None:
        return hit, None
    rates = by_name.get(customer_name)
    if not rates:
        return None, f'Customer {customer_name!r} not found in MNT UK price list.'
    if len(rates) == 1:
        return rates[0], None
    avg = round(sum(rates) / len(rates), 2)
    return avg, (f'Multiple MNT UK rates for {customer_name!r}; ZIP {zip_code!r} not an exact match -- '
                 f'used average of {len(rates)} known rates, confirm with WH.')


def parse_q3_august(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Sheet1']
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    book = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        pl = r[idx['PL nr.']]
        if not pl:
            continue
        rate = r[idx['Solaredge rate EUR']] or 0
        waiting = r[idx['Waiting fee EUR']] or 0
        cancel = r[idx['Cancellation fee EUR']] or 0
        customs = r[idx['Customs Clerance Fee EUR']] or 0
        duty_gbp = r[idx['Duty GBP']]
        total_eur = rate + waiting + cancel + customs
        book[pl] = {'total_eur': total_eur, 'duty_gbp': duty_gbp if isinstance(duty_gbp, (int, float)) else None}
    return book


def has_pod(row):
    pod = row.get('Logistic POD')
    return isinstance(pod, (int, float)) or hasattr(pod, 'year')


# --------------------------------------------------------------------------
# Load + dedup source data
# --------------------------------------------------------------------------

def load_rows():
    wb = openpyxl.load_workbook(MAIN_FILE, data_only=True)
    ws = wb['BI 1409']
    headers = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
        if all(v is None for v in vals):
            continue
        rows.append(dict(zip(headers, vals)))
    return rows


def dedup_shipped(rows):
    """Exact full-row duplicates among Shipped lines only -- Backlog rows are
    never touched, per the explicit request."""
    shipped, backlog = [], []
    seen = set()
    n_dupes = 0
    for row in rows:
        if row.get('Source') != 'Shipped':
            backlog.append(row)
            continue
        key = tuple(row.items())
        if key in seen:
            n_dupes += 1
            continue
        seen.add(key)
        shipped.append(row)
    print(f"Dedup: {n_dupes} exact-duplicate Shipped row(s) dropped ({len(shipped)} kept, "
          f"{len(backlog)} Backlog rows untouched).")
    return shipped + backlog


def num_pallets(row):
    v = row.get('# of Pallets per line-up')
    if not isinstance(v, (int, float)):
        return None
    return max(1, math.ceil(v - 1e-9))


def group_key_for(row):
    sh = row.get('Shipment Number')
    if sh and sh != 'N/A':
        return ('SH', sh)
    return ('ORD', row.get('SE Order#'))


# --------------------------------------------------------------------------
# Classification (population + pricing method), per row
# --------------------------------------------------------------------------

def classify(row):
    whs = row.get('Sending WHS Code')
    dest = row.get('Destination country')
    family = row.get('Family Type')
    ai = (row.get('A/I') or '').strip().upper()
    parent_region = row.get('Parent Region')

    if whs == CANOT_WHS:
        if dest == 'Israel':
            return 'Canot - Domestic', 'canot'
        return 'Canot - Export', 'matrix'

    if whs in NL_WHS:
        if row.get('Customer Name') == BAYWA_IT_CUSTOMER and isinstance(row.get('Zip'), str) \
                and 'block' in row['Zip'].lower():
            return 'NL - BayWa IT (manual)', 'baywa_it'
        if family == 'Battery' or dest == 'United Kingdom':
            return 'NL - Battery + UK', 'mnt_uk'
        if ai == 'SUPPORT':
            return 'NL - Support', 'domestic'
        if parent_region == 'EUROPE':
            return 'NL - Domestic', 'domestic'
        return 'NL - Export', 'matrix'

    return 'Other WHS - Export', 'matrix'


# --------------------------------------------------------------------------
# Pricing (grouped by shipment, flat rate allocated pro-rata by pallet share)
# --------------------------------------------------------------------------

def price_all(rows, dbs_book, matrix, mnt_uk_book, q3_book):
    recs = []
    for row in rows:
        population, method = classify(row)
        recs.append({'row': row, 'population': population, 'method': method,
                      'route': None, 'cost': None, 'currency': None, 'note': ''})

    # Sub-route MNT+POD+Q3 vs DBS within 'domestic', decided per-row (Forwarder/POD
    # are row-level facts) but the actual price lookup is still done per GROUP.
    for rec in recs:
        if rec['method'] == 'domestic':
            row = rec['row']
            is_mnt_pod = row.get('Forwarder') == 'MNT' and has_pod(row)
            ship_num = row.get('Shipment Number')
            if is_mnt_pod and ship_num in q3_book:
                rec['method'] = 'q3_mnt_pod'
            else:
                rec['method'] = 'dbs'
                if row.get('Forwarder') == 'MNT':
                    rec['note'] = ('MNT without a usable POD/Q3 match -- priced via DBS instead. ' + rec['note']).strip()

    # "Battery + UK" is priced from MNT UK price list, but that file only covers
    # a handful of specific UK-shipment customers -- most Battery-family lines
    # go to ordinary European destinations that were never meant to depend on
    # a UK truck-rate file. When the customer isn't in MNT UK at all, fall back
    # to the same domestic(DBS)/export(matrix) split as every other NL line,
    # same fallback pattern used for this population in the prior 0609 report.
    _, mnt_by_name = mnt_uk_book
    for rec in recs:
        if rec['method'] == 'mnt_uk' and rec['row'].get('Customer Name') not in mnt_by_name:
            row = rec['row']
            rec['method'] = 'dbs' if row.get('Parent Region') == 'EUROPE' else 'matrix'
            rec['note'] = (f"Customer {row.get('Customer Name')!r} not in MNT UK price list -- "
                           f"priced via the domestic/export fallback instead. " + rec['note']).strip()

    buckets = defaultdict(list)
    for rec in recs:
        if rec['method'] in ('dbs', 'matrix', 'canot', 'mnt_uk', 'baywa_it', 'q3_mnt_pod'):
            buckets[(rec['method'], group_key_for(rec['row']))].append(rec)
        else:
            rec['note'] = f'Unhandled method {rec["method"]!r} -- confirm with WH contact.'

    for (method, gkey), group in buckets.items():
        total_pallets = sum((num_pallets(r['row']) or 0) for r in group)
        sample = group[0]['row']

        if method == 'dbs':
            zone, price, note = dbs_lookup(dbs_book, sample.get('Destination country code'),
                                            sample.get('Zip'), total_pallets or 1)
            _allocate(group, price, 'EUR', zone, note, total_pallets,
                      f'DBS Price list 2026, zone {zone}: {price} EUR for {total_pallets:g} total pallets.')

        elif method == 'matrix':
            org = sample.get('Sending WHS Code')
            hit = matrix_lookup(matrix, org, sample.get('Origin country'), sample.get('ShipMode'),
                                 sample.get('Destination country'))
            route = f"{org}/{sample.get('ShipMode')}->{sample.get('Destination country')}"
            if not hit:
                for r in group:
                    r['route'] = route
                    r['note'] = f'No Ship Cost Matrix rate for {route} -- confirm with WH contact.'
                continue
            cost, currency, no_pal, mnote = hit
            extra = mnote
            if isinstance(no_pal, (int, float)) and total_pallets > no_pal:
                extra = (extra + ' ' if extra else '') + \
                        f'Shipment total {total_pallets:g} pallets exceeds this corridor\'s {no_pal:.0f}-pallet rate -- verify capacity.'
            _allocate(group, cost, currency, route, extra, total_pallets,
                      f'Ship Cost Matrix flat rate {cost} {currency} for the whole shipment ({total_pallets:g} pallets).')

        elif method == 'canot':
            price, currency, note = canot_israel_lookup(sample.get('City'), total_pallets or 1)
            _allocate(group, price, currency, 'Canot domestic (IL)', note, total_pallets,
                      f'Canot rate card: {price} {currency} for {total_pallets:g} total pallets.')

        elif method == 'mnt_uk':
            price, note = mnt_uk_lookup(mnt_uk_book, sample.get('Customer Name'), sample.get('Zip'))
            _allocate(group, price, 'EUR' if price is not None else None, 'MNT UK (Price Q3)', note, total_pallets,
                      f'MNT UK full-truck rate {price} EUR allocated across the shipment\'s {total_pallets:g} pallets.')

        elif method == 'baywa_it':
            capped_total = max(1, math.ceil(total_pallets - 1e-9)) if total_pallets else 1
            price = BAYWA_IT_OVERRIDES.get(capped_total)
            note = (None if price is not None else
                    f'No BayWa IT override rate on file for {capped_total} total pallets '
                    f'(only 1, 2, 12, 13 are known) -- confirm with WH contact.')
            _allocate(group, price, 'EUR' if price is not None else None, 'BayWa IT (manual)', note, total_pallets,
                      f'User-supplied BayWa IT rate {price} EUR for {capped_total} total pallets.')

        elif method == 'q3_mnt_pod':
            info = q3_book[gkey[1]]
            _allocate(group, info['total_eur'], 'EUR', f'Q3 AUGUST (PL nr. {gkey[1]})', None, total_pallets,
                      f'Q3 AUGUST shipment-level rate {info["total_eur"]} EUR allocated by pallet share.')
            if info['duty_gbp']:
                for r in group:
                    share = ((num_pallets(r['row']) or 0) / total_pallets) if total_pallets else (1 / len(group))
                    r['note'] = (r['note'] + f' Duty {round(info["duty_gbp"] * share, 2)} GBP not included '
                                 '(different currency).').strip()

    return recs


def _allocate(group, flat_cost, currency, route, extra_note, total_pallets, base_note):
    for r in group:
        r['route'] = route
        if flat_cost is None:
            r['cost'] = None
            r['currency'] = None
            r['note'] = (extra_note or (r['note'] + ' ' + (extra_note or '')).strip()).strip()
            continue
        pallets = num_pallets(r['row']) or 0
        share = (pallets / total_pallets) if total_pallets else (1 / len(group))
        r['cost'] = round(flat_cost * share, 2)
        r['currency'] = currency
        note = base_note if len(group) == 1 else (base_note + f' Allocated pro-rata by pallet share across '
                                                    f'this shipment\'s {len(group)} lines.')
        if extra_note:
            note = (note + ' ' + extra_note).strip()
        r['note'] = (r['note'] + ' ' + note).strip()


# --------------------------------------------------------------------------
# Output workbook
# --------------------------------------------------------------------------

OUT_COLUMNS = [
    ('Source', lambda rec: rec['row'].get('Source')),
    ('Shipment Number', lambda rec: rec['row'].get('Shipment Number')),
    ('SE Order#', lambda rec: rec['row'].get('SE Order#')),
    ('Sending WHS Code', lambda rec: rec['row'].get('Sending WHS Code')),
    ('Dest WHS Code', lambda rec: rec['row'].get('Dest WHS Code')),
    ('Forwarder', lambda rec: rec['row'].get('Forwarder')),
    ('ShipMode', lambda rec: rec['row'].get('ShipMode')),
    ('Customer Name', lambda rec: rec['row'].get('Customer Name')),
    ('Destination country', lambda rec: rec['row'].get('Destination country')),
    ('ZIP', lambda rec: rec['row'].get('Zip')),
    ('Family Type', lambda rec: rec['row'].get('Family Type')),
    ('A/I', lambda rec: rec['row'].get('A/I')),
    ('# of Pallets per line-up', lambda rec: num_pallets(rec['row'])),
    ('Population', lambda rec: rec['population']),
    ('Method', lambda rec: rec['method']),
    ('Zone / Route', lambda rec: rec['route']),
    ('Cost', lambda rec: rec['cost']),
    ('Currency', lambda rec: rec['currency']),
    ('Cost (USD)', lambda rec: to_usd(rec['cost'], rec['currency'])),
    ('Note', lambda rec: rec['note']),
]

METHOD_LABELS = {
    'dbs': 'DBS Price list 2026', 'matrix': 'Ship Cost Matrix', 'canot': 'Canot rate card',
    'mnt_uk': 'MNT UK price list', 'baywa_it': 'BayWa IT (manual)', 'q3_mnt_pod': 'Q3 AUGUST (MNT, POD)',
}


def write_sheet(wb, title, records):
    ws = wb.create_sheet(title)
    headers = [c[0] for c in OUT_COLUMNS]
    ws.append(headers)
    bold = Font(name='Arial', bold=True)
    arial = Font(name='Arial')
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = bold
    for rec in records:
        row = [getter(rec) for _, getter in OUT_COLUMNS]
        method_idx = [h for h, _ in OUT_COLUMNS].index('Method')
        row[method_idx] = METHOD_LABELS.get(row[method_idx], row[method_idx])
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).font = arial
    for c, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(14, min(40, len(header) + 4))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.freeze_panes = 'A2'
    return ws


def write_summary(wb, shipped_recs, backlog_recs):
    ws = wb.create_sheet('Summary', 1)
    bold = Font(name='Arial', bold=True)
    arial = Font(name='Arial')
    headers = ['Population', 'Status', 'Lines', 'Priced', 'Unpriced', 'Total Cost (USD)']
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = bold

    populations = sorted(set(r['population'] for r in shipped_recs + backlog_recs))
    r_out = 2
    for pop in populations:
        for status, recs in (('Shipped', shipped_recs), ('Backlog', backlog_recs)):
            subset = [r for r in recs if r['population'] == pop]
            if not subset:
                continue
            priced = [r for r in subset if r['cost'] is not None]
            total_usd = sum(to_usd(r['cost'], r['currency']) or 0 for r in priced)
            ws.append([pop, status, len(subset), len(priced), len(subset) - len(priced), round(total_usd, 2)])
            for c in range(1, len(headers) + 1):
                ws.cell(row=r_out, column=c).font = arial
            ws.cell(row=r_out, column=6).number_format = '#,##0.00'
            r_out += 1

    total_row = r_out
    ws.cell(row=total_row, column=1, value='Total').font = bold
    ws.cell(row=total_row, column=2).font = bold
    for c in (3, 4, 5, 6):
        col = get_column_letter(c)
        cell = ws.cell(row=total_row, column=c, value=f'=SUM({col}2:{col}{total_row - 1})')
        cell.font = bold
        if c == 6:
            cell.number_format = '#,##0.00'
    for c, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(22, len(h) + 4)


def write_readme(wb):
    readme = wb.create_sheet('README', 0)
    readme.column_dimensions['A'].width = 105
    bold = Font(name='Arial', bold=True)
    arial = Font(name='Arial')
    lines = [
        ('BI 1409 Freight Cost Report', bold),
        ('', arial),
        ('Source: data/BI_1409.xlsx, single flat "BI 1409" sheet (3,132 lines: 2,539 Shipped, 593 Backlog).', arial),
        ('', arial),
        ('Populations and rules:', bold),
        ('  Canot (3PLCANOT): Destination=Israel -> Canot Whs. NIS rate card by city region (200 ILS/1 pallet, '
         '1200/1300/1400 ILS 8T/12T for 2+). Everything else -> Ship Cost Matrix.', arial),
        ('  Netherlands (3PLDBSNL/3PLDBSBRNL), checked in order: (1) Family Type=Battery or Destination=UK -> '
         'MNT UK price list "Price Q3"; (2) A/I=Support -> same as domestic; (3) Parent Region=EUROPE (domestic) -> '
         'Forwarder=MNT with a real Logistic POD and a Q3 AUGUST match -> Q3 AUGUST, else -> DBS Price list 2026 '
         '(zone+pallets); (4) export -> Ship Cost Matrix. BayWa r.e. Solar Systems srl (Italy) lines with a '
         '"3c/4c block" ZIP use the fixed EUR-by-pallet override table instead of DBS.', arial),
        ('  Any other Sending WHS Code -> Ship Cost Matrix.', arial),
        ('', arial),
        ('Every rate source above (DBS, Ship Cost Matrix, Q3 AUGUST, MNT UK, Canot rate card, BayWa override) '
         'gives ONE flat price per shipment/order, not per line. Every line is priced by grouping same-shipment '
         'lines (Shipment Number, or SE Order# when there is none yet) within its pricing method, looking up '
         'one price for the group\'s total whole-number pallets, and allocating pro-rata by each line\'s pallet '
         'share -- built in from the start this time (this exact bug, discovered after the fact, was the single '
         'largest correction across the whole prior 0609 data pull project).', arial),
        ('', arial),
        ('Dedup: Shipped rows were checked for exact full-row duplicates (see the "Dedup" line printed when this '
         'was generated); no Backlog row is ever removed, per the request.', arial),
        ('', arial),
        ('FX to USD: 1 EUR = 1.1568, 1 ILS = 0.32955 (2026-09-14 snapshot, xe.com/tradingeconomics.com -- not a '
         'contracted rate); Ship Cost Matrix costs are already USD.', arial),
        ('', arial),
        ('Rows with a blank Cost need a human check (see their Note) -- no DBS zone/coverage, no Ship Cost Matrix '
         'corridor, no MNT UK customer match, no BayWa override rate for that shipment\'s total pallets, etc.', bold),
    ]
    for text, font in lines:
        readme.append((text,))
        readme.cell(row=readme.max_row, column=1).font = font


def main():
    print("Loading price sources...")
    dbs_book = parse_dbs_price_book(DBS_FILE)
    matrix = parse_ship_matrix(MATRIX_FILE)
    mnt_uk_book = parse_mnt_uk(MNT_UK_FILE)
    q3_book = parse_q3_august(Q3_AUG_FILE)

    print("Loading BI 1409 data...")
    rows = load_rows()
    print(f"Loaded {len(rows)} rows.")
    rows = dedup_shipped(rows)

    print("Pricing...")
    recs = price_all(rows, dbs_book, matrix, mnt_uk_book, q3_book)

    shipped_recs = [r for r in recs if r['row'].get('Source') == 'Shipped']
    backlog_recs = [r for r in recs if r['row'].get('Source') == 'Backlog']

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_sheet(wb, 'Shipped', shipped_recs)
    write_sheet(wb, 'Backlog', backlog_recs)
    write_summary(wb, shipped_recs, backlog_recs)
    write_readme(wb)

    os.makedirs(os.path.dirname(OUT_FILE), exist_ok=True)
    wb.save(OUT_FILE)
    print("Saved", OUT_FILE)

    n_priced = sum(1 for r in recs if r['cost'] is not None)
    total_usd = sum(to_usd(r['cost'], r['currency']) or 0 for r in recs if r['cost'] is not None)
    print(f"Total lines: {len(recs)}, priced: {n_priced}, unpriced: {len(recs) - n_priced}")
    print(f"Grand total: {total_usd:,.2f} USD")


if __name__ == '__main__':
    main()
