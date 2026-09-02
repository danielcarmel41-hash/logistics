"""
Same costing as scripts/nl_to_sup_cost.py, run against the new, much larger
data/NL_to_SUP_Data.xlsx (sheet "NL to SUP ", 2,699 lines: both regular
customer (SO) shipments and warehouse-transfer (TO) lines, all from the two
NL warehouses), split into two output sheets by "Source" as requested:
"Shipped" and "Backlog".

Costing (identical to the earlier NL to SUP run):
  - data/DBS_Price_list_2026.xlsx, zone = destination country code + first
    2 digits of ZIP (GB uses postcode-area letters; NL is a single zone),
    EP = "# of Pallets per line" rounded up to the next whole pallet
    bracket (1-33, 33 = Full load).
  - Cost (USD) = Cost (EUR) x 1.1618, the same FX snapshot used across this
    project (2026-09-01, xe.com/investing.com — not a contracted rate).
  - Manual override: BayWa r.e. Solar Systems srl (Italy) lines whose ZIP
    is garbage text ("3c/4c block") use the same user-supplied EUR rates
    by pallet count established earlier (1->240, 2->411, 12->2100,
    13->1410), since DBS can never match a non-numeric ZIP for them.

Columns match the ones requested for the earlier NL to SUP report:
Shipment Number, Dest WHS Code, Forwarder, Family Type, plus the
identifying/costing columns (Sending WHS Code, ZIP, Destination country,
Zone, # of Pallets per line, Cost EUR/USD, Note). "Dest WHS Code" is
"N/A" for the regular customer (SO) lines, which only carry a destination
country/customer — that's expected, not an error.

Output: output/NL_to_SUP_Cost.xlsx (replaces the earlier 43-row version),
sheets "Shipped" and "Backlog".
"""
import math
import os
import re

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = lambda name: os.path.join(BASE_DIR, "data", name)
OUT_FILE = os.path.join(BASE_DIR, "output", "NL_to_SUP_Cost.xlsx")

NL_SUP_FILE = DATA("NL_to_SUP_Data.xlsx")
DBS_FILE = DATA("DBS_Price_list_2026.xlsx")

MAX_EP = 33
EUR_TO_USD = 1.1618  # same snapshot used across this project (2026-09-01)

BAYWA_IT_CUSTOMER = 'BayWa r.e. Solar Systems srl'
BAYWA_IT_OVERRIDES = {1: 240, 2: 411, 12: 2100, 13: 1410}


def parse_dbs_price_book(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    book = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row = next(
            (r for r in range(1, 15)
             if any(ws.cell(row=r, column=c).value == 'EP' for c in range(1, ws.max_column + 1))),
            None,
        )
        if header_row is None:
            continue
        ep_col = None
        zone_cols = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if v == 'EP':
                ep_col = c
            elif isinstance(v, str) and v.strip().upper().startswith(sheet_name.upper()):
                zone_cols[v.strip()] = c
            elif isinstance(v, str) and sheet_name.upper() == 'NL' and re.match(r'^\d+\s*-\s*\d+$', v.strip()):
                zone_cols['NL'] = c

        zones = {z: {} for z in zone_cols}
        data_start = header_row + 2
        for r in range(data_start, ws.max_row + 1):
            ep_val = ws.cell(row=r, column=ep_col).value
            if not isinstance(ep_val, (int, float)):
                continue
            ep_val = int(ep_val)
            for zone, col in zone_cols.items():
                price = ws.cell(row=r, column=col).value
                if isinstance(price, (int, float)):
                    zones[zone][ep_val] = price
        book[sheet_name] = zones
    return book


# Postal codes are always exactly this many digits in these countries — a
# longer digit string starting with "0" (e.g. Swiss "04658" for Daniken,
# whose real PLZ is "4658") is a zero-padding artifact, not part of the
# code, and would otherwise resolve to the wrong zone (CH04 instead of
# CH46).
EXPECTED_ZIP_DIGITS = {'CH': 4, 'AT': 4, 'BE': 4, 'HU': 4, 'SI': 4, 'LU': 4, 'NL': 4,
                       'DE': 5, 'IT': 5, 'FR': 5, 'PL': 5, 'CZ': 5, 'FI': 5}


def dbs_zone_for(country_code, zip_code):
    country_code = (country_code or '').strip().upper()
    zip_code = str(zip_code or '').strip().upper()
    if country_code == 'NL':
        return 'NL'
    if country_code == 'GB':
        m = re.match(r'^[A-Z]+', zip_code.replace(' ', ''))
        return 'GB' + m.group() if m else None
    digits = ''.join(ch for ch in zip_code if ch.isdigit())
    expected = EXPECTED_ZIP_DIGITS.get(country_code)
    if expected and len(digits) == expected + 1 and digits[0] == '0':
        digits = digits[1:]
    if len(digits) < 2:
        return None
    return country_code + digits[:2]


def dbs_lookup(dbs_book, country_code, zip_code, pallets):
    zone = dbs_zone_for(country_code, zip_code)
    if not zone:
        return None, None, 'ZIP could not be parsed into a rate zone'
    zone_prices = dbs_book.get((country_code or '').strip().upper(), {}).get(zone)
    if not zone_prices:
        return zone, None, f'No DBS rate sheet/zone for country={country_code!r} zone={zone!r}'
    if not isinstance(pallets, (int, float)):
        return zone, None, f'Invalid pallet count on this line ({pallets!r})'
    ep = max(1, min(MAX_EP, math.ceil(pallets - 1e-9)))
    while ep <= MAX_EP:
        if ep in zone_prices:
            return zone, zone_prices[ep], None
        ep += 1
    return zone, None, f'Zone {zone} has no rate bracket for {pallets} pallets'


def load_rows():
    wb = openpyxl.load_workbook(NL_SUP_FILE, data_only=True)
    ws = wb['NL to SUP ']
    headers = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
        if all(v is None for v in vals):
            continue
        rows.append(dict(zip(headers, vals)))
    return rows


COLUMNS = [
    ('Shipment Number', 'shipment_number'),
    ('Dest WHS Code', 'dest_whs_code'),
    ('Forwarder', 'forwarder'),
    ('Family Type', 'family_type'),
    ('Sending WHS Code', 'sending_whs_code'),
    ('ZIP', 'zip'),
    ('Destination country', 'destination_country'),
    ('Zone', 'zone'),
    ('# of Pallets per line', 'pallets'),
    ('Cost (EUR)', 'cost_eur'),
    ('Cost (USD)', 'cost_usd'),
    ('Note', 'note'),
]


def price_row(row, dbs_book):
    pallets = row.get('# of Pallets per line')
    customer = row.get('Customer Name')
    zip_code = row.get('Zip')

    if customer == BAYWA_IT_CUSTOMER and isinstance(zip_code, str) and 'block' in zip_code.lower():
        cost_eur = BAYWA_IT_OVERRIDES.get(pallets)
        zone = 'BayWa IT (manual)'
        note = ('User-supplied rate for invalid ZIP.' if cost_eur is not None
                else f'No user-supplied rate for {pallets} pallets — please provide.')
    else:
        zone, cost_eur, note = dbs_lookup(
            dbs_book, row.get('Destination country code'), zip_code, pallets
        )
        note = note or ''

    cost_usd = round(cost_eur * EUR_TO_USD, 2) if cost_eur is not None else None
    return {
        'shipment_number': row.get('Shipment Number'),
        'dest_whs_code': row.get('Dest WHS Code'),
        'forwarder': row.get('Forwarder'),
        'family_type': row.get('Family Type'),
        'sending_whs_code': row.get('Sending WHS Code'),
        'zip': zip_code,
        'destination_country': row.get('Destination country'),
        'zone': zone,
        'pallets': pallets,
        'cost_eur': cost_eur,
        'cost_usd': cost_usd,
        'note': note,
    }


def write_sheet(wb, title, out_rows):
    ws = wb.create_sheet(title)
    headers = [c[0] for c in COLUMNS]
    ws.append(headers)
    bold = Font(name='Arial', bold=True)
    arial = Font(name='Arial')
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = bold
    for row in out_rows:
        ws.append([row[key] for _, key in COLUMNS])
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).font = arial
    for c, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(14, min(30, len(h) + 4))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.freeze_panes = 'A2'
    return ws


def main():
    dbs_book = parse_dbs_price_book(DBS_FILE)
    rows = load_rows()

    shipped_out = [price_row(r, dbs_book) for r in rows if r.get('Source') == 'Shipped']
    backlog_out = [price_row(r, dbs_book) for r in rows if r.get('Source') == 'Backlog']

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_sheet(wb, 'Shipped', shipped_out)
    write_sheet(wb, 'Backlog', backlog_out)

    readme = wb.create_sheet('README', 0)
    readme.column_dimensions['A'].width = 100
    bold = Font(name='Arial', bold=True)
    arial = Font(name='Arial')
    lines = [
        ('NL to SUP — pallet-based DBS cost, split by Shipped / Backlog', bold),
        ('', arial),
        ('Zone = destination country code + first 2 digits of ZIP (GB uses postcode-area letters; '
         'NL is a single zone); EP = # of Pallets per line, rounded up to the next whole bracket '
         '(1-33, 33 = Full load).', arial),
        ('Cost (USD) = Cost (EUR) x 1.1618, same FX snapshot used across this project (2026-09-01, '
         'xe.com/investing.com — not a contracted rate).', arial),
        ('BayWa r.e. Solar Systems srl (Italy) lines with a "3c/4c block" ZIP use the same '
         'user-supplied EUR rates by pallet count established earlier, not the DBS lookup.', arial),
        ('"Dest WHS Code" is N/A for regular customer (SO) shipments — only warehouse-transfer '
         '(TO) lines have a real one. That is expected, not an error.', arial),
    ]
    for text, font in lines:
        readme.append((text,))
        readme.cell(row=readme.max_row, column=1).font = font

    os.makedirs(os.path.dirname(OUT_FILE), exist_ok=True)
    wb.save(OUT_FILE)

    for label, out_rows in [('Shipped', shipped_out), ('Backlog', backlog_out)]:
        n_priced = sum(1 for r in out_rows if r['cost_eur'] is not None)
        print(f"{label}: {len(out_rows)} rows, {n_priced} priced, {len(out_rows) - n_priced} unpriced")


if __name__ == '__main__':
    main()
