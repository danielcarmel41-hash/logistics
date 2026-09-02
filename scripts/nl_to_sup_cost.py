"""
Costs every line on the "NL to SUP" sheet of data/NL_to_SUP_0209.xlsx (43
warehouse-transfer lines from the two NL warehouses to European supplier
warehouses) against data/DBS_Price_list_2026.xlsx, the same way as the
earlier Backlog/Shipped analyses:

  - each DBS sheet is one destination country; column G ("EP") holds the
    pallet-count bracket (1..33, 33 = "Full load")
  - the rate zone is destination country code + first 2 digits of the ZIP
    (single "NL" zone for Netherlands; not needed here since none of this
    file's destinations is the UK)
  - EP is looked up from this file's "# of Pallets per line " column,
    rounded UP to the next whole pallet (no pre-rounded column is provided
    here, unlike the earlier files)

Verified zones for this file's 3 destinations: NL (single zone), HU22
(Hungary, ZIP 2220), TR34 (Türkiye, ZIP 34500) — all present in the price
book.

Cost (USD) uses the same FX snapshot as the other outputs in this project
(1 EUR = 1.1618 USD, 2026-09-01, xe.com/investing.com — not a contracted
rate).

Output: output/NL_to_SUP_Cost.xlsx, one row per line with Shipment Number,
Dest WHS Code, Forwarder and Family Type as requested, plus the
identifying/costing columns.
"""
import math
import os
import re

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = lambda name: os.path.join(BASE_DIR, "data", name)
OUT_FILE = os.path.join(BASE_DIR, "output", "NL_to_SUP_Cost.xlsx")

NL_SUP_FILE = DATA("NL_to_SUP_0209.xlsx")
DBS_FILE = DATA("DBS_Price_list_2026.xlsx")

MAX_EP = 33
EUR_TO_USD = 1.1618  # same snapshot used across this project (2026-09-01)


def parse_dbs_price_book(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    book = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row = next(
            (r for r in range(1, 15)
             if any(ws.cell(row=r, column=c).value == 'EP' for c in range(1, ws.max_column + 1))),
            None,
        )
        if header_row is None:
            continue
        ep_col = None
        zone_cols = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if v == 'EP':
                ep_col = c
            elif isinstance(v, str) and v.strip().upper().startswith(sheet_name.upper()):
                zone_cols[v.strip()] = c
            elif isinstance(v, str) and sheet_name.upper() == 'NL' and re.match(r'^\d+\s*-\s*\d+$', v.strip()):
                zone_cols['NL'] = c

        zones = {z: {} for z in zone_cols}
        data_start = header_row + 2
        for r in range(data_start, ws.max_row + 1):
            ep_val = ws.cell(row=r, column=ep_col).value
            if not isinstance(ep_val, (int, float)):
                continue
            ep_val = int(ep_val)
            for zone, col in zone_cols.items():
                price = ws.cell(row=r, column=col).value
                if isinstance(price, (int, float)):
                    zones[zone][ep_val] = price
        book[sheet_name] = zones
    return book


def dbs_zone_for(country_code, zip_code):
    country_code = (country_code or '').strip().upper()
    zip_code = str(zip_code or '').strip().upper()
    if country_code == 'NL':
        return 'NL'
    if country_code == 'GB':
        m = re.match(r'^[A-Z]+', zip_code.replace(' ', ''))
        return 'GB' + m.group() if m else None
    digits = ''.join(ch for ch in zip_code if ch.isdigit())
    if len(digits) < 2:
        return None
    return country_code + digits[:2]


def dbs_lookup(dbs_book, country_code, zip_code, pallets):
    zone = dbs_zone_for(country_code, zip_code)
    if not zone:
        return None, None, 'ZIP could not be parsed into a rate zone'
    zone_prices = dbs_book.get((country_code or '').strip().upper(), {}).get(zone)
    if not zone_prices:
        return zone, None, f'No DBS rate sheet/zone for country={country_code!r} zone={zone!r}'
    if not isinstance(pallets, (int, float)):
        return zone, None, f'Invalid pallet count on this line ({pallets!r})'
    ep = max(1, min(MAX_EP, math.ceil(pallets - 1e-9)))
    while ep <= MAX_EP:
        if ep in zone_prices:
            return zone, zone_prices[ep], None
        ep += 1
    return zone, None, f'Zone {zone} has no rate bracket for {pallets} pallets'


def main():
    dbs_book = parse_dbs_price_book(DBS_FILE)

    wb = openpyxl.load_workbook(NL_SUP_FILE, data_only=True)
    ws = wb['NL to SUP']
    headers = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
        if all(v is None for v in vals):
            continue
        rows.append(dict(zip(headers, vals)))

    out_rows = []
    for row in rows:
        pallets = row.get('# of Pallets per line ')
        zone, cost_eur, note = dbs_lookup(
            dbs_book, row.get('Destination country code'), row.get('Zip'), pallets
        )
        cost_usd = round(cost_eur * EUR_TO_USD, 2) if cost_eur is not None else None
        out_rows.append({
            'shipment_number': row.get('Shipment Number'),
            'dest_whs_code': row.get('Dest WHS Code'),
            'forwarder': row.get('Forwarder'),
            'family_type': row.get('Family Type'),
            'sending_whs_code': row.get('Sending WHS Code'),
            'zip': row.get('Zip'),
            'destination_country': row.get('Destination country'),
            'zone': zone,
            'pallets': pallets,
            'cost_eur': cost_eur,
            'cost_usd': cost_usd,
            'note': note or '',
        })

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = 'NL to SUP Cost'
    columns = [
        ('Shipment Number', 'shipment_number'),
        ('Dest WHS Code', 'dest_whs_code'),
        ('Forwarder', 'forwarder'),
        ('Family Type', 'family_type'),
        ('Sending WHS Code', 'sending_whs_code'),
        ('ZIP', 'zip'),
        ('Destination country', 'destination_country'),
        ('Zone', 'zone'),
        ('# of Pallets per line', 'pallets'),
        ('Cost (EUR)', 'cost_eur'),
        ('Cost (USD)', 'cost_usd'),
        ('Note', 'note'),
    ]
    headers_out = [c[0] for c in columns]
    ws_out.append(headers_out)
    bold = Font(name='Arial', bold=True)
    arial = Font(name='Arial')
    for c in range(1, len(headers_out) + 1):
        ws_out.cell(row=1, column=c).font = bold
    for row in out_rows:
        ws_out.append([row[key] for _, key in columns])
    for r in range(2, ws_out.max_row + 1):
        for c in range(1, len(headers_out) + 1):
            ws_out.cell(row=r, column=c).font = arial
    for c, h in enumerate(headers_out, start=1):
        ws_out.column_dimensions[get_column_letter(c)].width = max(14, min(30, len(h) + 4))
    ws_out.auto_filter.ref = f"A1:{get_column_letter(len(headers_out))}{ws_out.max_row}"
    ws_out.freeze_panes = 'A2'

    readme = wb_out.create_sheet('README', 0)
    readme.column_dimensions['A'].width = 100
    lines = [
        ('NL to SUP — pallet-based DBS cost, from the "NL to SUP" sheet', bold),
        ('', arial),
        ('Zone = destination country code + first 2 digits of ZIP (NL is a single zone); '
         'EP = # of Pallets per line, rounded up to the next whole pallet bracket (1-33, '
         '33 = Full load).', arial),
        ('Cost (USD) = Cost (EUR) x 1.1618, same FX snapshot used across this project '
         '(2026-09-01, xe.com/investing.com — not a contracted rate).', arial),
    ]
    for text, font in lines:
        readme.append((text,))
        readme.cell(row=readme.max_row, column=1).font = font

    os.makedirs(os.path.dirname(OUT_FILE), exist_ok=True)
    wb_out.save(OUT_FILE)

    n_priced = sum(1 for r in out_rows if r['cost_eur'] is not None)
    print(f"Rows: {len(out_rows)}, priced: {n_priced}")
    for r in out_rows:
        if r['cost_eur'] is None:
            print("  UNPRICED:", r)


if __name__ == '__main__':
    main()
