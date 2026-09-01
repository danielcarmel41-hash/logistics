"""
Sums the Q3 AUGUST freight cost per customer for the UK + Battery lines that
were pulled out of the main Shipped file (data/UK_Battery_Data.xlsx — 146
Backlog lines, 23 distinct customers).

These lines have no Shipment Number yet (Backlog), so they can't be joined
to data/Q3_prices_AUGUST.xlsx by "PL nr." the way the Shipped/MNT lines were
in the earlier analysis. Per this request, the join is by customer name
instead: for every customer in UK_Battery_Data, every row in Q3_prices_
AUGUST.xlsx whose "Ship to " matches that customer is summed on column A
("Solaredge rate EUR") only (no waiting/cancellation/customs fees this
time), then converted to USD with the same FX snapshot used earlier
(1 EUR = 1.1618 USD, see scripts/pallet_cost_truck_sea.py).

Matching: exact name match first; a few customers are the same company
under a slightly different legal-suffix spelling between the two files
(e.g. "Krannich Solar GmbH & Co. KG" here vs "Krannich Solar GmbH" in Q3),
so a fuzzy fallback matches when one normalized name is a word-for-word
prefix of the other. This is careful not to blur genuinely different
entities such as "Krannich Solar AG" vs "Krannich Solar GmbH" (kept apart
because "AG" and "GmbH" differ at the same word position). Customers with
no Q3 AUGUST data at all (they didn't ship via MNT in August) are listed
with a zero/blank total and flagged.

No Shipment Number / "PL nr." matching is used anywhere in this script —
these Backlog lines don't have one yet, and the join is by customer name
only, per this request.

Output: output/UK_Battery_Cost_Summary.xlsx
  - "Summary"     one row per UK_Battery_Data customer: matched Q3 name,
                   number of Q3 shipments summed, total EUR, total USD.
  - "Q3 Detail"    every Q3 row that fed into a customer's total (by name),
                   for verification.
"""
import os
import re

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = lambda name: os.path.join(BASE_DIR, "data", name)
OUT_FILE = os.path.join(BASE_DIR, "output", "UK_Battery_Cost_Summary.xlsx")

UKB_FILE = DATA("UK_Battery_Data.xlsx")
Q3_FILE = DATA("Q3_prices_AUGUST.xlsx")

EUR_TO_USD = 1.1618  # same snapshot as scripts/pallet_cost_truck_sea.py (2026-09-01)


def norm(name):
    s = (name or '').upper()
    s = s.replace('&', ' ').replace('.', ' ').replace(',', ' ')
    return re.sub(r'\s+', ' ', s).strip()


def prefix_match(a, b):
    wa, wb = a.split(), b.split()
    n = min(len(wa), len(wb))
    return n > 0 and wa[:n] == wb[:n]


def load_ukb_customers():
    wb = openpyxl.load_workbook(UKB_FILE, data_only=True)
    ws = wb['Sheet2']
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    customers = []
    seen = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        name = r[idx['Customer Name']]
        if name and name not in seen:
            seen.add(name)
            customers.append(name)
    return sorted(customers)


def load_q3_rows():
    wb = openpyxl.load_workbook(Q3_FILE, data_only=True)
    ws = wb['Sheet1']
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        ship_to = r[idx['Ship to ']]
        rate = r[idx['Solaredge rate EUR']]
        if not ship_to:
            continue
        rows.append({
            'ship_to': ship_to,
            'rate_eur': rate if isinstance(rate, (int, float)) else None,
            'pl_nr': r[idx['PL nr.']],
            'country': r[idx['country']],
            'nr_pal': r[idx['Nr. of pal']],
        })
    return rows


def match_customer(ukb_customer, q3_names_norm):
    nc = norm(ukb_customer)
    exact = [q for q in q3_names_norm if q == nc]
    if exact:
        return exact, 'exact'
    fuzzy = [q for q in q3_names_norm if prefix_match(q, nc)]
    return fuzzy, 'fuzzy' if fuzzy else 'none'


def main():
    ukb_customers = load_ukb_customers()
    q3_rows = load_q3_rows()
    q3_by_norm_name = {}
    for row in q3_rows:
        q3_by_norm_name.setdefault(norm(row['ship_to']), []).append(row)

    summary_rows = []
    detail_rows = []

    for customer in ukb_customers:
        matched_norm_names, match_type = match_customer(customer, q3_by_norm_name.keys())
        matched_q3_rows = []
        for nn in matched_norm_names:
            matched_q3_rows.extend(q3_by_norm_name[nn])

        priced = [r for r in matched_q3_rows if r['rate_eur'] is not None]
        total_eur = sum(r['rate_eur'] for r in priced)
        total_usd = round(total_eur * EUR_TO_USD, 2) if priced else None

        matched_q3_display_names = sorted({r['ship_to'] for r in matched_q3_rows})
        note = ''
        if match_type == 'none':
            note = 'No Q3 AUGUST shipments found for this customer name — total is blank.'
        elif match_type == 'fuzzy':
            note = (f'Matched by name similarity to {", ".join(matched_q3_display_names)} '
                     '(exact spelling differs) — verify this is the same customer.')
        if matched_q3_rows and not priced:
            note = (note + ' ' if note else '') + 'Matched shipments exist but have no rate in column A.'

        summary_rows.append({
            'customer': customer,
            'matched_q3_name': ', '.join(matched_q3_display_names) if matched_q3_display_names else None,
            'n_shipments': len(priced),
            'total_eur': round(total_eur, 2) if priced else None,
            'total_usd': total_usd,
            'note': note,
        })

        for r in matched_q3_rows:
            detail_rows.append({
                'customer': customer,
                'q3_ship_to': r['ship_to'],
                'country': r['country'],
                'nr_pal': r['nr_pal'],
                'rate_eur': r['rate_eur'],
                'rate_usd': round(r['rate_eur'] * EUR_TO_USD, 2) if isinstance(r['rate_eur'], (int, float)) else None,
            })

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold = Font(name='Arial', bold=True)
    arial = Font(name='Arial')

    ws = wb.create_sheet('Summary')
    headers = ['Customer Name (UK/Battery file)', 'Matched Q3 Customer Name', 'Q3 Shipments Summed',
               'Total Solaredge Rate (EUR)', 'Total Cost (USD)', 'Note']
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = bold
    for row in summary_rows:
        ws.append([row['customer'], row['matched_q3_name'], row['n_shipments'],
                   row['total_eur'], row['total_usd'], row['note']])
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).font = arial
    for c, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(16, min(45, len(h) + 4))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.freeze_panes = 'A2'

    ws2 = wb.create_sheet('Q3 Detail')
    headers2 = ['Customer Name (UK/Battery file)', 'Q3 Ship to', 'Country', 'Nr. of pal',
                'Solaredge rate (EUR)', 'Solaredge rate (USD)']
    ws2.append(headers2)
    for c in range(1, len(headers2) + 1):
        ws2.cell(row=1, column=c).font = bold
    for row in detail_rows:
        ws2.append([row['customer'], row['q3_ship_to'], row['country'], row['nr_pal'],
                    row['rate_eur'], row['rate_usd']])
    for r in range(2, ws2.max_row + 1):
        for c in range(1, len(headers2) + 1):
            ws2.cell(row=r, column=c).font = arial
    for c, h in enumerate(headers2, start=1):
        ws2.column_dimensions[get_column_letter(c)].width = max(16, min(40, len(h) + 4))
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers2))}{ws2.max_row}"
    ws2.freeze_panes = 'A2'

    readme = wb.create_sheet('README', 0)
    readme.column_dimensions['A'].width = 100
    lines = [
        ('UK / Battery — Q3 AUGUST cost summary by customer', bold),
        ('', arial),
        ('These 146 lines (23 customers) are Backlog, so they have no Shipment Number yet and '
         'cannot be joined to Q3_prices_AUGUST by "PL nr." — matched by Customer Name instead.', arial),
        ('Total per customer = SUM of "Solaredge rate EUR" (column A) over every Q3 AUGUST row for '
         'that customer, converted to USD at 1 EUR = 1.1618 USD (same snapshot used in '
         'Pallet_Cost_by_Ship_Method.xlsx, 2026-09-01, xe.com/investing.com — not a contracted rate).', arial),
        ('19 of 23 customers matched (16 exact, 3 by a legal-suffix spelling difference — see the Note '
         'column). 4 customers have no Q3 AUGUST shipments at all: Energia Italia S.P.A. a Socio Unico, '
         'European Solar Technology Group b.v., SAM HYDRO-SOLAR, VP Solar srl — their total is blank.', arial),
    ]
    for text, font in lines:
        readme.append((text,))
        readme.cell(row=readme.max_row, column=1).font = font

    os.makedirs(os.path.dirname(OUT_FILE), exist_ok=True)
    wb.save(OUT_FILE)

    matched = sum(1 for r in summary_rows if r['total_usd'] is not None)
    print(f"Customers: {len(summary_rows)}, with a USD total: {matched}")
    for row in summary_rows:
        print(f"  {row['customer']:45} {row['total_eur']!s:>10} EUR  {row['total_usd']!s:>10} USD  ({row['note']})")


if __name__ == '__main__':
    main()
