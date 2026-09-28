"""
Regenerates output/BI_2809_Freight_Cost_Report.xlsx from source with the
updated bi_2809_freight_cost_report rules (relaxed Battery+UK/Domestic/
Support estimates, NL - Support + SEA -> Ship Cost Matrix, BayWa IT
reclassified into NL - Domestic, DEFAULT-forwarder lines treated as DBS),
then overlays the user's own edits from their uploaded, hand-edited copy
(data/BI_2809_Freight_Cost_Report_user_edit.xlsx) so that nothing in the
four populations they manually corrected is disturbed:

    NL - Export (SEA), Canot - Export, Other WHS - Export,
    Other WHS - Export backlog

These four populations are untouched by every rule change in this round
(they are all priced via 'matrix'/'canot', neither of which changed), so
a fresh regeneration reproduces byte-identical Cost/Currency/Note values
for them -- except for the exact rows the user has since hand-edited
(zeroed a cost) or deleted outright. This script detects those by row key
(Source, Shipment Number, SE Order#, Customer Name, Population, Method)
against the user's uploaded file and:
  - copies the user's Cost/Currency/Cost (USD)/Note back onto the matching
    row in the fresh output (so a manually-zeroed row stays zero), and
  - deletes any fresh row in one of these four populations that has no
    matching row in the user's file at all (a row they removed).
No row outside these four populations is touched by this step -- every
other population's rows come straight from the freshly recomputed report.

The Summary sheet is rebuilt afterwards from the final Shipped/Backlog/
Manual Review sheets so its totals reflect both the rule changes and the
preserved user edits.
"""
import os
import sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(BASE_DIR, 'scripts'))
import bi_2809_freight_cost_report as core

USER_EDIT_FILE = os.path.join(BASE_DIR, 'data', 'BI_2809_Freight_Cost_Report_user_edit.xlsx')
OUT_FILE = core.OUT_FILE

PROTECTED_POPULATIONS = {
    'NL - Export (SEA)', 'Canot - Export', 'Other WHS - Export', 'Other WHS - Export backlog',
}

OVERRIDE_FIELDS = ('Cost', 'Currency', 'Cost (USD)', 'Note')


def row_key(get):
    return (get('Source'), get('Shipment Number'), get('SE Order#'), get('Customer Name'),
            get('Population'), get('Method'))


def load_protected_rows(wb, sheet_name):
    """key -> row dict, for the protected populations only, from the user's
    uploaded edited file."""
    ws = wb[sheet_name]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in r):
            continue
        d = dict(zip(headers, r))
        if d.get('Population') not in PROTECTED_POPULATIONS:
            continue
        get = lambda h: d.get(h)
        out[row_key(get)] = d
    return out


def apply_protected_overrides(out_wb, edit_wb):
    for sheet_name in ('Shipped', 'Backlog'):
        edit_rows = load_protected_rows(edit_wb, sheet_name)
        ws = out_wb[sheet_name]
        headers = [c.value for c in ws[1]]
        col = {h: i + 1 for i, h in enumerate(headers)}

        rows_to_delete = []
        matched = 0
        for r in range(2, ws.max_row + 1):
            get = lambda h: ws.cell(row=r, column=col[h]).value
            if get('Population') not in PROTECTED_POPULATIONS:
                continue
            key = row_key(get)
            erow = edit_rows.pop(key, None)
            if erow is None:
                rows_to_delete.append(r)
                continue
            for field in OVERRIDE_FIELDS:
                ws.cell(row=r, column=col[field]).value = erow.get(field)
            matched += 1

        for r in reversed(rows_to_delete):
            ws.delete_rows(r)

        print(f"  {sheet_name}: preserved {matched} protected-population row(s) from the user's edit, "
              f"removed {len(rows_to_delete)} row(s) the user deleted.")
        if edit_rows:
            print(f"  WARNING: {len(edit_rows)} protected row(s) in the user's edit had no match in the "
                  f"fresh output for {sheet_name} -- not applied: {list(edit_rows.keys())[:5]}")

        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def rebuild_summary(wb, manual_sheet_name='Manual Review'):
    if 'Summary' in wb.sheetnames:
        del wb['Summary']
    ws = wb.create_sheet('Summary', 1)
    bold = Font(name='Arial', bold=True)
    arial = Font(name='Arial')
    headers = ['Population', 'Status', 'Lines (consolidated rows)', 'Priced', 'Unpriced', 'Total Cost (USD)']
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = bold

    by_pop_status = defaultdict(lambda: [0, 0, 0.0])
    for sheet_name in ('Shipped', 'Backlog'):
        ws_data = wb[sheet_name]
        idx = {c.value: i for i, c in enumerate(ws_data[1])}
        for r in ws_data.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in r):
                continue
            key = (r[idx['Population']], sheet_name)
            agg = by_pop_status[key]
            agg[0] += 1
            cost_usd = r[idx['Cost (USD)']]
            if cost_usd is not None:
                agg[1] += 1
                agg[2] += cost_usd

    populations = sorted(set(k[0] for k in by_pop_status))
    r_out = 2
    for pop in populations:
        for status in ('Shipped', 'Backlog'):
            agg = by_pop_status.get((pop, status))
            if not agg:
                continue
            ws.append([pop, status, agg[0], agg[1], agg[0] - agg[1], round(agg[2], 2)])
            for c in range(1, len(headers) + 1):
                ws.cell(row=r_out, column=c).font = arial
            ws.cell(row=r_out, column=6).number_format = '#,##0.00'
            r_out += 1

    manual_ws = wb[manual_sheet_name]
    n_manual = manual_ws.max_row - 1 if manual_ws.max_row > 1 else 0
    ws.append(['Manual Review (not priced)', 'Mixed', n_manual, 0, n_manual, 0])
    for c in range(1, len(headers) + 1):
        ws.cell(row=r_out, column=c).font = arial
    r_out += 1

    total_row = r_out
    ws.cell(row=total_row, column=1, value='Total').font = bold
    ws.cell(row=total_row, column=2).font = bold
    for c in (3, 4, 5, 6):
        col_letter = get_column_letter(c)
        cell = ws.cell(row=total_row, column=c, value=f'=SUM({col_letter}2:{col_letter}{total_row - 1})')
        cell.font = bold
        if c == 6:
            cell.number_format = '#,##0.00'
    for c, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(24, len(h) + 4)


def update_readme(wb):
    ws = wb['README']
    bold = Font(name='Arial', bold=True)
    arial = Font(name='Arial')
    ws.append((None,))
    ws.append(('Patched per user review (this round):',))
    ws.cell(row=ws.max_row, column=1).font = bold
    lines = [
        '1) "NL - Battery + UK", "NL - Domestic" and "NL - Support" MNT/Q3-priced groups that had no exact '
        'Shipment Number/POD match are now given a price ESTIMATE from Q3 AUGUST: the nearest available pallet '
        'count on file for the same destination country, preferring the same customer name when more than one '
        'is on file. "NL - Domestic"/"NL - Support" DBS-priced groups whose ZIP could not be resolved to a zone '
        '(or whose zone has no bracket for that many pallets) are similarly estimated from the DBS Price list at '
        'the country level (averaged across every zone on file for that country at the same pallet count), '
        'ignoring ZIP. Every such row is clearly flagged "Estimated ..." in its Note for confirmation with WH.',
        '2) "NL - Support" lines shipped by SEA are now priced from the general Ship Cost Matrix (by origin/'
        'destination country), instead of the DBSCHENKER/MNT forwarder split used for LAND Support lines.',
        '3) "NL - BayWa IT (manual)" is retired as its own population now that it has real prices -- those lines '
        'are reclassified into "NL - Domestic" (their Method column still reads "BayWa IT (manual)").',
        '4) A blank/"DEFAULT" Forwarder value (no real forwarder on file) is now treated the same as DBSCHENKER '
        'by default -- confirmed against several named examples (MARCHIOL S.P.A, Ecostal Yomatec, Shipment '
        'Number SH10826976919) that were previously sitting unpriced in "Manual Review". A genuinely different, '
        'named forwarder (DGF, Q4, DSV, ...) is unaffected and still left in "Manual Review" for manual pricing.',
        '5) The Cost/Currency/Cost (USD)/Note of every row in "NL - Export (SEA)", "Canot - Export", "Other WHS '
        '- Export" and "Other WHS - Export backlog" -- the four populations the user hand-priced/zeroed in their '
        'own review -- are taken verbatim from that reviewed file, including one row they removed outright; '
        'none of this round\'s rule changes touch these four populations anyway (all priced via Ship Cost '
        'Matrix/Canot rate card, neither of which changed).',
    ]
    for text in lines:
        ws.append((text,))
        ws.cell(row=ws.max_row, column=1).font = arial


def main():
    print("Regenerating the full report from source with the updated rules...")
    core.main()

    print("Applying the user's protected-population edits on top...")
    edit_wb = openpyxl.load_workbook(USER_EDIT_FILE, data_only=True)
    out_wb = openpyxl.load_workbook(OUT_FILE)

    apply_protected_overrides(out_wb, edit_wb)
    rebuild_summary(out_wb)
    update_readme(out_wb)

    out_wb.save(OUT_FILE)
    print("Saved", OUT_FILE)

    ws_summary = out_wb['Summary']
    total_row = ws_summary.max_row
    print(f"Summary total row: {[ws_summary.cell(row=total_row, column=c).value for c in range(1, 7)]}")


if __name__ == '__main__':
    main()
