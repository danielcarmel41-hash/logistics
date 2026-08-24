"""
Prices every shipment line in data/DB_NL.xlsx using the per-country,
per-zone, per-pallet-count rate tables in data/DBS_Price_list_2026.xlsx.

Rate lookup logic
------------------
1. Each sheet in the price list is a destination country (e.g. "DE").
   Row 9 holds the rate zone codes as column headers (e.g. DE01, DE04, ...).
   Column G ("EP") holds the pallet count for each rate row (1..33, where
   33 is the "Full load" bracket). A price lives at the intersection.
2. The rate zone for a shipment is derived from its destination ZIP:
     - Most countries: <country code> + first two digits of the ZIP
       (e.g. Germany 73776 -> DE73). This matches the "2 Digit zipcodes"
       label the price list itself uses for the Dutch sheet.
     - Great Britain: the ZIP's leading postcode-area letters
       (e.g. "LE3 1BY" -> GBLE).
     - Netherlands: the price list only has a single zone column, so
       every NL shipment maps to it regardless of ZIP.
3. The pallet count is rounded up to the nearest rate bracket (1..33);
   anything above 33 pallets is billed at the "Full load" (33) rate.

Output: output/DB_NL_with_shipping_cost.xlsx with the requested columns
(Sending WHS Code, ZIP, Customer Name, # of Pallets per line) plus the
matched Zone, Shipping Cost and Family Type, with an active AutoFilter
over the whole table. Rows whose ZIP could not be matched to a zone (bad
data in the source file, e.g. a ZIP field containing "3c block" instead
of a postal code) are flagged in the Note column; a small number of these
(the Italy "3c/4c block" lines) have a manually supplied cost in
MANUAL_COST_OVERRIDES below, provided by the user from the actual invoice.
"""
import os
import re

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PRICE_FILE = os.path.join(BASE_DIR, "data", "DBS_Price_list_2026.xlsx")
DATA_FILE = os.path.join(BASE_DIR, "data", "DB_NL.xlsx")
OUT_FILE = os.path.join(BASE_DIR, "output", "DB_NL_with_shipping_cost.xlsx")

MAX_EP = 33  # last row in every price sheet = "Full load" bracket

# Manual cost overrides (EUR) for shipment lines whose ZIP could not be
# resolved to a rate zone automatically. Keyed by the row number in the
# "DB" sheet of data/DB_NL.xlsx. Supplied by the user from the actual
# invoiced cost for these BayWa (Italy) "3c/4c block" lines.
MANUAL_COST_OVERRIDES = {
    78: 2100,
    84: 240,
    753: 1410,
    754: 411,
    755: 240,
    756: 240,
    757: 240,
    758: 240,
    759: 240,
    760: 411,
    761: 411,
    762: 411,
    763: 240,
    764: 240,
    765: 240,
    766: 240,
    767: 240,
    768: 240,
    769: 240,
    770: 240,
    771: 240,
    772: 411,
}


def parse_price_book(path):
    """Returns {country_code: {zone_code: {pallet_count: price}}}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    book = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row = next(
            (r for r in range(1, 15)
             if any(ws.cell(row=r, column=c).value == 'EP' for c in range(1, ws.max_column + 1))),
            None,
        )
        if header_row is None:
            continue

        ep_col = None
        zone_cols = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if v == 'EP':
                ep_col = c
            elif isinstance(v, str) and v.strip().upper().startswith(sheet_name.upper()):
                zone_cols[v.strip()] = c
            elif isinstance(v, str) and sheet_name.upper() == 'NL' and re.match(r'^\d+\s*-\s*\d+$', v.strip()):
                # NL has a single zone column labelled "10 - 99" instead of an "NLxx" code
                zone_cols['NL'] = c

        zones = {z: {} for z in zone_cols}
        data_start = header_row + 2  # the header row is duplicated one row below
        for r in range(data_start, ws.max_row + 1):
            ep_val = ws.cell(row=r, column=ep_col).value
            if not isinstance(ep_val, (int, float)):
                continue
            ep_val = int(ep_val)
            for zone, col in zone_cols.items():
                price = ws.cell(row=r, column=col).value
                if isinstance(price, (int, float)):
                    zones[zone][ep_val] = price
        book[sheet_name] = zones
    return book


def zone_for(country_code, zip_code):
    country_code = (country_code or '').strip().upper()
    zip_code = str(zip_code or '').strip().upper()

    if country_code == 'NL':
        return 'NL'

    if country_code == 'GB':
        m = re.match(r'^[A-Z]+', zip_code.replace(' ', ''))
        return 'GB' + m.group() if m else None

    digits = ''.join(ch for ch in zip_code if ch.isdigit())
    if len(digits) < 2:
        return None
    return country_code + digits[:2]


def price_for_pallets(zone_prices, pallets):
    if pallets is None:
        return None
    ep = max(1, min(MAX_EP, int(round(pallets + 0.4999))))  # round up, cap at full-load bracket
    while ep <= MAX_EP:
        if ep in zone_prices:
            return zone_prices[ep]
        ep += 1
    return None


def main():
    price_book = parse_price_book(PRICE_FILE)

    wb2 = openpyxl.load_workbook(DATA_FILE, data_only=True)
    ws2 = wb2['DB']
    headers = [c.value for c in ws2[1]]
    idx = {h: i + 1 for i, h in enumerate(headers)}

    col_whs = idx['Sending WHS Code']
    col_zip = idx['Zip']
    col_cust = idx['Customer Name']
    col_pallets = idx['# of Pallets per line']
    col_country_code = idx['Destination country code']
    col_country = idx['Destination country']
    col_family_type = idx['Family Type']

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = 'Shipping Cost'

    out_headers = [
        'Sending WHS Code', 'ZIP', 'Customer Name', '# of Pallets per line',
        'Destination country', 'Zone', 'Shipping Cost (EUR)', 'Family Type', 'Note',
    ]
    out_ws.append(out_headers)
    for c in range(1, len(out_headers) + 1):
        out_ws.cell(row=1, column=c).font = Font(name='Arial', bold=True)

    unmatched = []
    row_out = 2
    for r in range(2, ws2.max_row + 1):
        whs = ws2.cell(row=r, column=col_whs).value
        zip_code = ws2.cell(row=r, column=col_zip).value
        cust = ws2.cell(row=r, column=col_cust).value
        pallets = ws2.cell(row=r, column=col_pallets).value
        country = ws2.cell(row=r, column=col_country).value
        country_code = ws2.cell(row=r, column=col_country_code).value
        family_type = ws2.cell(row=r, column=col_family_type).value

        if whs is None and zip_code is None and cust is None:
            continue

        zone = zone_for(country_code, zip_code)
        zone_prices = price_book.get(country_code, {}).get(zone) if zone else None
        price = price_for_pallets(zone_prices, pallets) if zone_prices else None

        note = ''
        if price is None:
            if r in MANUAL_COST_OVERRIDES:
                price = MANUAL_COST_OVERRIDES[r]
                note = 'Manually supplied cost (ZIP is not a valid postal code)'
            else:
                note = 'ZIP could not be matched to a rate zone (check source ZIP value)'
                unmatched.append((r, whs, zip_code, cust, pallets, country, country_code))

        out_ws.cell(row=row_out, column=1, value=whs)
        out_ws.cell(row=row_out, column=2, value=zip_code)
        out_ws.cell(row=row_out, column=3, value=cust)
        out_ws.cell(row=row_out, column=4, value=pallets)
        out_ws.cell(row=row_out, column=5, value=country)
        out_ws.cell(row=row_out, column=6, value=zone)
        out_ws.cell(row=row_out, column=7, value=price)
        out_ws.cell(row=row_out, column=8, value=family_type)
        out_ws.cell(row=row_out, column=9, value=note)
        row_out += 1

    for c, header in enumerate(out_headers, start=1):
        out_ws.column_dimensions[get_column_letter(c)].width = max(14, len(header) + 2)

    arial_font = Font(name='Arial')
    for row in out_ws.iter_rows(min_row=2, max_row=out_ws.max_row):
        for cell in row:
            cell.font = arial_font

    last_col_letter = get_column_letter(len(out_headers))
    out_ws.auto_filter.ref = f"A1:{last_col_letter}{row_out - 1}"

    os.makedirs(os.path.dirname(OUT_FILE), exist_ok=True)
    out_wb.save(OUT_FILE)

    print(f"Rows written: {row_out - 2}")
    print(f"Unmatched rows: {len(unmatched)}")
    for u in unmatched:
        print(u)


if __name__ == '__main__':
    main()
